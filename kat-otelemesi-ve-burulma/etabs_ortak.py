# -*- coding: utf-8 -*-
"""
etabs_ortak.py
--------------
burulma_kontrol.py ve drift_kontrol.py tarafından kullanılan ortak fonksiyonlar.
ÜÇ DOSYA AYNI KLASÖRDE OLMALIDIR.

Yaklaşım
  Kat ötelemeleri ETABS'ın "Story Drifts" tablosundan DEĞİL, nokta yer
  değiştirmelerinden hesaplanır. Üst kattaki her nokta için alt kat yer
  değiştirmesi şöyle alınır:
    * alt katta aynı (x, y) konumunda nokta varsa -> doğrudan o noktadan,
    * yoksa (plan küçülmesi / büyümesi, aks kayması, çıkma) -> alt katın
      rijit cisim hareketinden (ux, uy, θz; en küçük kareler uyumu).
  Böylece her katın uç noktaları O KATIN KENDİ PLANINDAN belirlenir; plan
  küçüldüğünde alt katın artık var olmayan uç noktaları kullanılmaz.

Gereksinimler:  pip install comtypes numpy pandas openpyxl
"""

# ---------------------------------------------------------------------------
# SORUMLULUK REDDİ
# Bu araç bir mühendislik yardımcısıdır. Hesap, tasarım ve kontrol sorumluluğu
# tamamen kullanıcıya aittir. Çıktılar, ilgili yönetmelik ve proje koşullarına
# göre kullanıcı tarafından doğrulanmadan hiçbir projede kullanılmamalıdır.
# Yazar, kullanımdan doğabilecek hiçbir zarardan sorumlu tutulamaz.
# ---------------------------------------------------------------------------
from __future__ import annotations

import difflib
import math
from dataclasses import dataclass, field

import numpy as np
import pandas as pd

# ---------------------------------------------------------------------------
# Toleranslar (gerekirse değiştirin)
# ---------------------------------------------------------------------------
XY_TOL = 0.05        # m   iki noktanın alt/üst katta "aynı aks" sayılma toleransı
Z_TOL = 0.02         # m   nokta kotu ile kat kotu arasındaki izin verilen fark
RIJIT_SAPMA = 0.05   # -   rijit diyafram uyum hatası / en büyük yer değiştirme
MIN_YAYILIM = 1.0    # m   burulma için uç noktalar arası en küçük mesafe

SEVIYE_SIRA = {"HATA": 0, "UYARI": 1, "BİLGİ": 2}


# ---------------------------------------------------------------------------
# Uyarı kaydı
# ---------------------------------------------------------------------------
class Uyarilar:
    def __init__(self):
        self._k: list[dict] = []

    def ekle(self, seviye: str, mesaj: str, durum: str = "", kat: str = ""):
        self._k.append({"Seviye": seviye, "Durum": durum, "Kat": kat, "Mesaj": mesaj})

    def tablo(self) -> pd.DataFrame:
        df = pd.DataFrame(self._k, columns=["Seviye", "Durum", "Kat", "Mesaj"]).drop_duplicates()
        if df.empty:
            return df
        df["_s"] = df["Seviye"].map(SEVIYE_SIRA)
        return df.sort_values(["_s", "Durum"], kind="stable").drop(columns="_s")

    def say(self, seviye: str) -> int:
        return int((self.tablo()["Seviye"] == seviye).sum()) if self._k else 0


# ---------------------------------------------------------------------------
# ETABS bağlantısı ve veri okuma
# ---------------------------------------------------------------------------
def _kontrol(r, ne: str):
    ret = r[-1] if isinstance(r, (list, tuple)) else r
    if ret != 0:
        raise RuntimeError(f"ETABS API çağrısı başarısız: {ne} (dönüş kodu {ret})")
    return r


def etabs_baglan():
    try:
        import comtypes.client
    except ImportError:
        raise SystemExit("comtypes yüklü değil. Komut satırında:  pip install comtypes")
    try:
        etabs = comtypes.client.GetActiveObject("CSI.ETABS.API.ETABSObject")
    except OSError:
        raise SystemExit("Çalışan bir ETABS bulunamadı. ETABS'ı açın, modeli yükleyin ve analizi çalıştırın.")
    model = etabs.SapModel
    model.SetPresentUnits(6)  # kN_m_C  -> tüm uzunluklar metre
    if not model.GetModelIsLocked():
        raise SystemExit("Model kilitli değil, yani analiz sonucu yok. Önce analizi çalıştırın (F5).")
    return model


def katlari_oku(model) -> pd.DataFrame:
    """Kat adları ve kotları, tabandan yukarı sıralı. İlk satır 'Base'."""
    try:
        r = _kontrol(model.Story.GetStories_2(), "Story.GetStories_2")
        taban, adlar, kotlar = float(r[0]), list(r[2]), list(r[3])
    except Exception:
        r = _kontrol(model.Story.GetStories(), "Story.GetStories")
        adlar, kotlar, yuk = list(r[1]), list(r[2]), list(r[3])
        i = int(np.argmin(kotlar))
        taban = kotlar[i] - yuk[i]
    kat = pd.DataFrame({"Kat": adlar, "Kot": kotlar})
    if not (kat["Kat"].str.lower() == "base").any():
        kat = pd.concat([pd.DataFrame({"Kat": ["Base"], "Kot": [taban]}), kat])
    return kat.sort_values("Kot").reset_index(drop=True)


def noktalari_oku(model, grup: str) -> pd.DataFrame:
    if grup.lower() == "all":
        r = _kontrol(model.PointObj.GetAllPoints(), "PointObj.GetAllPoints")
        adlar, xs, ys, zs = list(r[1]), list(r[2]), list(r[3]), list(r[4])
    else:
        r = _kontrol(model.GroupDef.GetAssignments(grup), f"GroupDef.GetAssignments('{grup}')")
        adlar = [ad for tip, ad in zip(r[1], r[2]) if tip == 1]  # 1 = nokta
        if not adlar:
            raise SystemExit(f"'{grup}' grubunda hiç nokta yok.")
        xs, ys, zs = [], [], []
        for ad in adlar:
            c = _kontrol(model.PointObj.GetCoordCartesian(ad), "PointObj.GetCoordCartesian")
            xs.append(c[0]); ys.append(c[1]); zs.append(c[2])

    kayit = []
    for ad, x, y, z in zip(adlar, xs, ys, zs):
        etiket, kat = "", ""
        lab = model.PointObj.GetLabelFromName(ad)
        if lab[-1] == 0:
            etiket, kat = lab[0], lab[1]
        a = b = c = 0.0
        try:
            la = model.PointObj.GetLocalAxes(ad)
            if la[-1] == 0:
                a, b, c = float(la[0]), float(la[1]), float(la[2])
        except Exception:
            pass
        diy = ""
        try:
            d = model.PointObj.GetDiaphragm(ad)
            if d[-1] == 0:
                diy = {1: "Bağlantısız", 2: "Döşemeden"}.get(int(d[0]), str(d[1]))
        except Exception:
            pass
        kayit.append(dict(Nokta=str(ad), Etiket=etiket, Kat=kat, X=x, Y=y, Z=z,
                          EksenA=a, EksenB=b, EksenC=c, Diyafram=diy))
    return pd.DataFrame(kayit)


def noktalari_hazirla(noktalar: pd.DataFrame, katlar: pd.DataFrame, uyari: Uyarilar) -> pd.DataFrame:
    """Kat numarası atar; kat kotunda olmayan, bağlantısız, ikiz noktaları ayıklar."""
    df = noktalar.copy()
    if "Diyafram" in df:
        bag = df["Diyafram"] == "Bağlantısız"
        if bag.any():
            uyari.ekle("BİLGİ", f"{int(bag.sum())} nokta diyaframdan bağlantısız olarak tanımlı; hesap dışı.")
        df = df[~bag]

    kat_no = {k.lower(): i for i, k in enumerate(katlar["Kat"])}
    df["KatNo"] = df["Kat"].astype(str).str.lower().map(kat_no)
    yok = df["KatNo"].isna()
    if yok.any():
        uyari.ekle("UYARI", f"{int(yok.sum())} noktanın kat etiketi tanınmadı; hesap dışı "
                            f"(örn. {', '.join(df.loc[yok, 'Nokta'].head(5))}).")
    df = df[~yok].copy()
    df["KatNo"] = df["KatNo"].astype(int)

    kot = katlar["Kot"].to_numpy()
    ara = (df["Z"] - kot[df["KatNo"].to_numpy()]).abs() > Z_TOL
    for kat, g in df[ara].groupby("Kat"):
        uyari.ekle("UYARI", f"{len(g)} nokta kat kotunda değil (ara kot, eğimli eleman, kiriş/kolon ara "
                            f"noktası olabilir) ve hesap dışı bırakıldı: {', '.join(g['Nokta'].head(6))}"
                            + (" ..." if len(g) > 6 else ""), kat=kat)
    df = df[~ara]

    if {"EksenB", "EksenC"} <= set(df.columns):
        egik = (df["EksenB"].abs() > 1e-6) | (df["EksenC"].abs() > 1e-6)
        if egik.any():
            uyari.ekle("UYARI", f"{int(egik.sum())} noktanın lokal ekseni X/Y dışında döndürülmüş; "
                                "sadece Z etrafındaki dönüş düzeltildi.")
    else:
        df["EksenA"] = df.get("EksenA", 0.0)

    anahtar = df["KatNo"].astype(str) + "_" + (df["X"] / XY_TOL).round().astype(int).astype(str) \
        + "_" + (df["Y"] / XY_TOL).round().astype(int).astype(str)
    ikiz = anahtar.duplicated()
    if ikiz.any():
        uyari.ekle("UYARI", f"{int(ikiz.sum())} nokta aynı katta başka bir noktayla çakışıyor (ikiz nokta); "
                            f"ilki kullanıldı: {', '.join(df.loc[ikiz, 'Nokta'].head(6))}")
    return df[~ikiz].reset_index(drop=True)


def _ad_listesi(r) -> list[str]:
    return [str(a) for a in r[1]] if r[-1] == 0 and r[0] > 0 else []


def _durum_tipi(model, ad):
    for f in ("GetTypeOAPI_1", "GetTypeOAPI"):
        try:
            r = getattr(model.LoadCases, f)(ad)
            if r[-1] == 0:
                return int(r[0])
        except Exception:
            continue
    return None


def _kombo_isaretsiz(model, ad, gorulen) -> bool:
    """Kombinasyon zarf/SRSS/mutlak ise veya içinde spektral durum varsa True."""
    if ad in gorulen:
        return False
    gorulen.add(ad)
    try:
        t = model.RespCombo.GetTypeOAPI(ad)
        if t[-1] == 0 and int(t[0]) in (1, 2, 3, 4):  # zarf, mutlak, SRSS, aralık
            return True
        r = model.RespCombo.GetCaseList(ad)
    except Exception:
        return False
    if r[-1] != 0:
        return False
    for tip, alt in zip(r[1], r[2]):
        if int(tip) == 1:
            if _kombo_isaretsiz(model, alt, gorulen):
                return True
        elif _durum_tipi(model, alt) in (4, 5, 6):  # spektral / zaman tanım alanı
            return True
    return False


def durum_bilgisi(model, isim: str):
    """(gerçek ad, 'Kombinasyon'|'Yükleme durumu', işaretsiz_mi)"""
    kombolar = _ad_listesi(model.RespCombo.GetNameList())
    durumlar = _ad_listesi(model.LoadCases.GetNameList())
    k_map = {k.lower(): k for k in kombolar}
    d_map = {d.lower(): d for d in durumlar}
    if isim.lower() in k_map:
        ad = k_map[isim.lower()]
        return ad, "Kombinasyon", _kombo_isaretsiz(model, ad, set())
    if isim.lower() in d_map:
        ad = d_map[isim.lower()]
        tip = _durum_tipi(model, ad)
        if tip == 3:
            raise ValueError(f"'{ad}' modal bir durum; mod şekilleri öteleme kontrolü için kullanılamaz.")
        return ad, "Yükleme durumu", tip in (4, 5, 6)
    oneri = difflib.get_close_matches(isim, kombolar + durumlar, n=5, cutoff=0.3)
    raise ValueError(f"'{isim}' modelde yükleme durumu veya kombinasyon olarak bulunamadı."
                     + (f" Bunlardan biri mi: {', '.join(oneri)}" if oneri else ""))


def yer_degistirmeleri_oku(model, isim: str, tur: str, grup: str) -> pd.DataFrame:
    s = model.Results.Setup
    s.DeselectAllCasesAndCombosForOutput()
    if tur == "Kombinasyon":
        _kontrol(s.SetComboSelectedForOutput(isim), f"SetComboSelectedForOutput('{isim}')")
    else:
        _kontrol(s.SetCaseSelectedForOutput(isim), f"SetCaseSelectedForOutput('{isim}')")
    for f, v in (("SetOptionMultiStepStatic", 2), ("SetOptionMultiValuedCombo", 2)):
        try:
            getattr(s, f)(v)  # çok adımlı sonuçları adım adım al
        except Exception:
            pass
    r = model.Results.JointDispl(grup, 2)  # 2 = grup
    if r[-1] != 0 or r[0] == 0:
        raise RuntimeError(f"'{isim}' için yer değiştirme sonucu alınamadı. Durum analiz edilmemiş "
                           "(Set Load Cases to Run) veya grupta sonuç üreten nokta yok.")
    return pd.DataFrame({"Nokta": [str(a) for a in r[1]], "AdimTipi": list(r[4]),
                         "AdimNo": list(r[5]), "U1": list(r[6]), "U2": list(r[7])})


# ---------------------------------------------------------------------------
# Hesap çekirdeği (ETABS'tan bağımsız; test edilebilir)
# ---------------------------------------------------------------------------
def rijit_uyum(x, y, ux, uy):
    """Kat noktalarına rijit cisim hareketi (u0, v0, θ) uydurur. sapma: bağıl RMS hata."""
    x, y, ux, uy = map(lambda a: np.asarray(a, float), (x, y, ux, uy))
    n = len(x)
    if n == 0:
        return None
    xc, yc = float(x.mean()), float(y.mean())
    if n == 1 or (np.ptp(x) < MIN_YAYILIM and np.ptp(y) < MIN_YAYILIM):
        return dict(xc=xc, yc=yc, u0=float(ux.mean()), v0=float(uy.mean()), th=0.0, sapma=0.0)
    A = np.zeros((2 * n, 3))
    A[:n, 0] = 1.0
    A[:n, 2] = -(y - yc)
    A[n:, 1] = 1.0
    A[n:, 2] = x - xc
    b = np.concatenate([ux, uy])
    p, *_ = np.linalg.lstsq(A, b, rcond=None)
    res = b - A @ p
    olcek = max(float(np.max(np.abs(b))), 1e-12)
    return dict(xc=xc, yc=yc, u0=float(p[0]), v0=float(p[1]), th=float(p[2]),
                sapma=float(np.sqrt(np.mean(res ** 2)) / olcek))


SIFIR_HAREKET = dict(xc=0.0, yc=0.0, u0=0.0, v0=0.0, th=0.0, sapma=0.0)


def rijit_hareket(u, x, y):
    x, y = np.asarray(x, float), np.asarray(y, float)
    return u["u0"] - u["th"] * (y - u["yc"]), u["v0"] + u["th"] * (x - u["xc"])


def plan_alani(x, y) -> float:
    p = sorted(set(zip(np.round(np.asarray(x, float), 3), np.round(np.asarray(y, float), 3))))
    if len(p) < 3:
        return 0.0

    def cr(o, a, b):
        return (a[0] - o[0]) * (b[1] - o[1]) - (a[1] - o[1]) * (b[0] - o[0])

    alt, ust = [], []
    for q in p:
        while len(alt) >= 2 and cr(alt[-2], alt[-1], q) <= 0:
            alt.pop()
        alt.append(q)
    for q in reversed(p):
        while len(ust) >= 2 and cr(ust[-2], ust[-1], q) <= 0:
            ust.pop()
        ust.append(q)
    h = alt[:-1] + ust[:-1]
    return 0.5 * abs(sum(h[i][0] * h[(i + 1) % len(h)][1] - h[(i + 1) % len(h)][0] * h[i][1]
                         for i in range(len(h))))


def _eta(D):
    dmax, dmin = float(np.max(D)), float(np.min(D))
    dort = (dmax + dmin) / 2.0
    eta = dmax / dort if dort > 0.01 * abs(dmax) and dmax > 0 else float("nan")
    return dmax, dmin, dort, eta


def kat_otelemeleri(katlar, noktalar, disp, yon, uyari: Uyarilar, etiket: str):
    """
    katlar  : Kat, Kot (tabandan yukarı, ilk satır Base)
    noktalar: Nokta, Etiket, Kat, KatNo, X, Y
    disp    : Nokta, UX, UY  (global, tek sonuç seti)
    Döner   : (kat_tablosu, nokta_tablosu)
    """
    d = noktalar.merge(disp, on="Nokta", how="left")
    eksik = d["UX"].isna()
    if eksik.any():
        uyari.ekle("UYARI", f"{int(eksik.sum())} noktada sonuç yok, hesaba katılmadı "
                            f"(örn. {', '.join(d.loc[eksik, 'Nokta'].head(5))}).", etiket)
    d = d[~eksik]
    dik_eksen = "Y" if yon == "X" else "X"
    kat_satir, nokta_satir = [], []
    top_ana = top_dik = 0.0

    for i in range(1, len(katlar)):
        kat = katlar.at[i, "Kat"]
        U = d[d["KatNo"] == i]
        if U.empty:
            uyari.ekle("BİLGİ", "Katta kontrol noktası yok, atlandı.", etiket, kat)
            continue
        j = i - 1
        while j > 0 and d[d["KatNo"] == j].empty:
            j -= 1
        L = d[d["KatNo"] == j]
        alt_kat = katlar.at[j, "Kat"]
        h = float(katlar.at[i, "Kot"] - katlar.at[j, "Kot"])
        if j != i - 1:
            uyari.ekle("UYARI", f"{katlar.at[i - 1, 'Kat']} katında nokta yok; öteleme {alt_kat} katına "
                                f"göre h = {h:.2f} m ile hesaplandı.", etiket, kat)
        if h <= 0:
            uyari.ekle("HATA", f"Kat yüksekliği {h:.3f} m; kat tanımlarını kontrol edin.", etiket, kat)
            continue

        alt_uyum = rijit_uyum(L["X"], L["Y"], L["UX"], L["UY"]) if len(L) else dict(SIFIR_HAREKET)
        ust_uyum = rijit_uyum(U["X"], U["Y"], U["UX"], U["UY"])

        Lx, Ly = L["X"].to_numpy(), L["Y"].to_numpy()
        Lux, Luy = L["UX"].to_numpy(), L["UY"].to_numpy()
        if len(L):
            kx0, kx1, ky0, ky1 = Lx.min() - XY_TOL, Lx.max() + XY_TOL, Ly.min() - XY_TOL, Ly.max() + XY_TOL
        a_ux, a_uy, kaynak = np.zeros(len(U)), np.zeros(len(U)), []
        for k, (x, y) in enumerate(zip(U["X"], U["Y"])):
            if len(L):
                mes = np.hypot(Lx - x, Ly - y)
                m = int(np.argmin(mes))
                if mes[m] <= XY_TOL:
                    a_ux[k], a_uy[k] = Lux[m], Luy[m]
                    kaynak.append("Alt kat noktası")
                    continue
            a_ux[k], a_uy[k] = rijit_hareket(alt_uyum, x, y)
            if not len(L):
                kaynak.append("Taban (sıfır)")
            elif not (kx0 <= x <= kx1 and ky0 <= y <= ky1):
                kaynak.append("Alt kat rijit hareketi (plan dışı)")
            else:
                kaynak.append("Alt kat rijit hareketi")

        dux = U["UX"].to_numpy() - a_ux
        duy = U["UY"].to_numpy() - a_uy
        D, Dd = (dux, duy) if yon == "X" else (duy, dux)
        isaret = 1.0 if D.mean() >= 0 else -1.0
        D = D * isaret

        dmax, dmin, dort, eta = _eta(D)
        yayilim = float(np.ptp(U[dik_eksen]))
        if np.max(np.abs(D)) < 1e-7:      # < 0.0001 mm: ötelenme yok (rijit bodrum vb.)
            eta = float("nan")
        top_ana += float(np.sum(np.abs(D)))
        top_dik += float(np.sum(np.abs(Dd)))
        if dmin < 0 < dmax:
            uyari.ekle("UYARI", "Bazı noktalar ters yönde ötelenmiş (burulma baskın). η formülü anlamını "
                                "yitirir; perde yerleşimini gözden geçirin.", etiket, kat)
        if yayilim < MIN_YAYILIM:
            uyari.ekle("UYARI", f"{dik_eksen} doğrultusunda uç noktalar arası mesafe {yayilim:.2f} m; "
                                "η hesaplanamadı.", etiket, kat)
            eta = float("nan")

        eta_r = float("nan")
        if ust_uyum is not None and len(U) > 1:
            ru = rijit_hareket(ust_uyum, U["X"], U["Y"])
            ra = rijit_hareket(alt_uyum, U["X"], U["Y"])
            Dr = (ru[0] - ra[0]) if yon == "X" else (ru[1] - ra[1])
            eta_r = _eta(Dr * isaret)[3]
        if ust_uyum["sapma"] > RIJIT_SAPMA:
            uyari.ekle("UYARI", f"Kat rijit diyafram gibi davranmıyor (uyum sapması %{100 * ust_uyum['sapma']:.1f}): "
                                "yarı-rijit diyafram, dilatasyonla ayrılmış bloklar veya diyafram atanmamış "
                                "noktalar olabilir. Uç nokta değerleri kullanıldı.", etiket, kat)

        n_ara = sum(k.startswith("Alt kat rijit") for k in kaynak)
        n_dis = sum("plan dışı" in k for k in kaynak)
        if n_ara:
            uyari.ekle("BİLGİ", f"{n_ara} noktanın {alt_kat} katında karşılığı yok; alt katın rijit "
                                "hareketinden hesaplandı" + (f" ({n_dis} tanesi alt kat planının dışında: "
                                "çıkma/konsol)." if n_dis else "."), etiket, kat)
            if alt_uyum["sapma"] > RIJIT_SAPMA:
                uyari.ekle("UYARI", f"{alt_kat} katı rijit davranmadığı için karşılıksız noktalardaki "
                                    "öteleme yaklaşıktır.", etiket, kat)

        A_u = plan_alani(U["X"], U["Y"])
        A_l = plan_alani(L["X"], L["Y"]) if len(L) >= 3 else 0.0
        oran = A_u / A_l if A_l > 0 else float("nan")
        plan = ""
        if oran == oran:
            if oran < 0.90:
                plan = f"Küçülme (%{100 * oran:.0f})"
            elif oran > 1.10:
                plan = f"Büyüme (%{100 * oran:.0f})"
        if plan:
            uyari.ekle("BİLGİ", f"Plan değişimi: {plan}. Uç noktalar bu katın kendi planından alındı.", etiket, kat)

        imax, imin, iabs = int(np.argmax(D)), int(np.argmin(D)), int(np.argmax(np.abs(D)))
        Un = U.reset_index(drop=True)
        kat_satir.append({
            "Kat": kat, "Alt kat": alt_kat, "h (m)": h, "Nokta sayısı": len(U),
            "Δmax (mm)": 1000 * dmax, "Δmax noktası": Un.at[imax, "Nokta"],
            "Δmin (mm)": 1000 * dmin, "Δmin noktası": Un.at[imin, "Nokta"],
            "Δort (mm)": 1000 * dort, "η": eta, "η (rijit kontrol)": eta_r,
            "|Δ|max (mm)": 1000 * abs(D[iabs]), "|Δ|max noktası": Un.at[iabs, "Nokta"],
            "|Δ|max X": Un.at[iabs, "X"], "|Δ|max Y": Un.at[iabs, "Y"],
            "Rijit sapma (%)": 100 * ust_uyum["sapma"], "Plan alanı (m²)": A_u, "Plan değişimi": plan,
            "Alt katta karşılıksız nokta": n_ara,
        })
        for k in range(len(Un)):
            nokta_satir.append({
                "Kat": kat, "Nokta": Un.at[k, "Nokta"], "Etiket": Un.at[k, "Etiket"],
                "X": Un.at[k, "X"], "Y": Un.at[k, "Y"],
                f"Δ{yon} (mm)": 1000 * D[k], f"Δ{dik_eksen} (mm)": 1000 * Dd[k],
                "Alt kat değeri": kaynak[k],
            })
    kdf = pd.DataFrame(kat_satir)
    kdf.attrs["yanlis_yon"] = top_dik > top_ana
    kdf.attrs["dik_oran"] = top_dik / top_ana if top_ana > 0 else float("inf")
    if not kdf.empty:
        genel = kdf["|Δ|max (mm)"].max()
        kdf["Öteleme ~0"] = kdf["|Δ|max (mm)"] <= max(1e-4, 1e-3 * genel)
        if genel <= 1e-4:
            uyari.ekle("HATA", "Hiçbir katta bu doğrultuda öteleme yok: durum adı veya yön (X/Y) yanlış "
                               "girilmiş olabilir.", etiket)
        elif kdf["Öteleme ~0"].any():
            uyari.ekle("BİLGİ", "Şu katlarda öteleme yok denecek kadar küçük (rijit bodrum / toprak perdesi "
                                "kuşatması beklenen durum): " + ", ".join(kdf.loc[kdf["Öteleme ~0"], "Kat"]), etiket)
    return kdf, pd.DataFrame(nokta_satir)


@dataclass
class DurumSonucu:
    isim: str
    yon: str
    tur: str
    isaretsiz: bool
    setler: list = field(default_factory=list)  # [(set etiketi, kat_df, nokta_df)]
    yanlis_yon: bool = False


def hamdan_hesapla(katlar, noktalar, ham, isim, yon, tur, isaretsiz, uyari) -> DurumSonucu:
    ham = ham.merge(noktalar[["Nokta", "EksenA"]], on="Nokta", how="inner")
    a = np.radians(ham["EksenA"].to_numpy())
    ham["UX"] = ham["U1"] * np.cos(a) - ham["U2"] * np.sin(a)
    ham["UY"] = ham["U1"] * np.sin(a) + ham["U2"] * np.cos(a)
    ham["AdimTipi"] = ham["AdimTipi"].fillna("").astype(str)
    gruplar = list(ham.groupby(["AdimTipi", "AdimNo"], sort=False))
    sonuc = DurumSonucu(isim, yon, tur, isaretsiz)
    for (tip, no), g in gruplar:
        et = "" if len(gruplar) == 1 else f"{tip} {no:g}".strip()
        g = g.drop_duplicates("Nokta")
        kdf, ndf = kat_otelemeleri(katlar, noktalar, g[["Nokta", "UX", "UY"]], yon, uyari,
                                   f"{isim} {et}".strip())
        sonuc.setler.append((et, kdf, ndf))
    sonuc.yanlis_yon = bool(sonuc.setler) and all(k.attrs.get("yanlis_yon", False) for _, k, _ in sonuc.setler)
    if sonuc.yanlis_yon:
        oran = min(k.attrs["dik_oran"] for _, k, _ in sonuc.setler)
        diger = "Y" if yon == "X" else "X"
        uyari.ekle("HATA", f"Bu durum {yon} listesine girilmiş ama bina {diger} doğrultusunda ötelenmiş "
                           f"({diger}/{yon} öteleme oranı {oran:.1f}). Yanlış listeye yazılmış olabilir; "
                           "sonuçlar değerlendirmeye alınmadı.", isim)
    if len(gruplar) > 1:
        uyari.ekle("BİLGİ", f"{len(gruplar)} sonuç seti var (çok adımlı statik veya zarf); her biri ayrı "
                            "hesaplandı ve en elverişsizi alındı.", isim)
    return sonuc


def durumu_hesapla(model, katlar, noktalar, grup, isim, yon, uyari) -> DurumSonucu:
    ad, tur, isaretsiz = durum_bilgisi(model, isim)
    ham = yer_degistirmeleri_oku(model, ad, tur, grup)
    return hamdan_hesapla(katlar, noktalar, ham, ad, yon, tur, isaretsiz, uyari)


# ---------------------------------------------------------------------------
# Excel çıktısı
# ---------------------------------------------------------------------------
def excel_yaz(dosya: str, sayfalar: dict) -> str:
    from openpyxl.styles import Font, PatternFill

    renk = {"yesil": "C6EFCE", "turuncu": "FFEB9C", "kirmizi": "FFC7CE", "gri": "E7E6E6"}

    def boya(deger: str):
        s = str(deger)
        if s in ("UYGUN", "YOK") or s == "BİLGİ":
            return renk["yesil"]
        if s.startswith("A1b") or s == "UYARI":
            return renk["turuncu"]
        if "AŞIL" in s or s in ("HATA", "η > 2.0") or "HESAPLANAMADI" in s:
            return renk["kirmizi"]
        if s in ("KONTROL DIŞI", "ÖTELEME YOK"):
            return renk["gri"]
        return None

    def yaz(hedef):
        with pd.ExcelWriter(hedef, engine="openpyxl") as w:
            for ad, df in sayfalar.items():
                ad = ad[:31]
                df.to_excel(w, sheet_name=ad, index=False)
                ws = w.sheets[ad]
                ws.freeze_panes = "A2"
                for hucre in ws[1]:
                    hucre.font = Font(bold=True)
                for ci, sutun in enumerate(df.columns, start=1):
                    genislik = max([len(str(sutun))] + [len(f"{v:.3f}" if isinstance(v, float) else str(v))
                                                         for v in df[sutun].head(200)])
                    ws.column_dimensions[ws.cell(1, ci).column_letter].width = min(genislik + 2, 90)
                    if sutun in ("Sonuç", "Seviye"):
                        for ri, v in enumerate(df[sutun], start=2):
                            r = boya(v)
                            if r:
                                ws.cell(ri, ci).fill = PatternFill("solid", fgColor=r)
                    elif df[sutun].dtype.kind == "f":
                        for ri in range(2, len(df) + 2):
                            ws.cell(ri, ci).number_format = "0.000"
    try:
        yaz(dosya)
        return dosya
    except PermissionError:
        from datetime import datetime
        yeni = dosya.replace(".xlsx", f"_{datetime.now():%H%M%S}.xlsx")
        yaz(yeni)
        return yeni


def model_verisini_hazirla(grup: str):
    uyari = Uyarilar()
    model = etabs_baglan()
    katlar = katlari_oku(model)
    noktalar = noktalari_hazirla(noktalari_oku(model, grup), katlar, uyari)
    print(f"{len(katlar) - 1} kat, {len(noktalar)} kontrol noktası okundu.")
    return model, katlar, noktalar, uyari
